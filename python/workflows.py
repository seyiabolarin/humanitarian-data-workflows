"""Offline, synthetic-data adaptations of Seyi Abolarin's register workflows."""
from pathlib import Path
from html import escape
import csv
import hashlib
import io
from PIL import Image, ImageOps
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as PDFImage

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs"


def load_register(path=None):
    """Preserve identifiers as text; reject ambiguous or incomplete exports."""
    path = Path(path or ROOT / "data/synthetic/register.csv")
    if path.suffix.lower() == ".xlsx":
        import pandas as pd
        rows = pd.read_excel(path, dtype=str).fillna("").to_dict("records")
    else:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
    required = {"record_id", "name", "community", "lga", "household_size"}
    if not rows or not required.issubset(rows[0]):
        raise ValueError("Supply a nonempty register with the documented columns.")
    ids = [r["record_id"].strip() for r in rows]
    if any(not x for x in ids) or len(set(ids)) != len(ids):
        raise ValueError("Record IDs must be nonempty and unique; review duplicates first.")
    return sorted(rows, key=lambda r: (r["lga"].casefold(), r["community"].casefold(), r["name"].casefold(), r["record_id"]))


def clean_image(source, destination, max_size=(640, 640)):
    """Apply EXIF orientation and remove metadata without modifying the source."""
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("Source and output image must be different files.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as opened:
        oriented = ImageOps.exif_transpose(opened).convert("RGB")
        oriented.thumbnail(max_size)
        clean = Image.new("RGB", oriented.size)
        clean.paste(oriented)
        clean.save(destination, format="JPEG", quality=82)
    return destination


def compress_folder(source, destination):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if source == destination or source in destination.parents or destination in source.parents:
        raise ValueError("Input and output folders must be separate, non-nested folders.")
    outputs = []
    for path in sorted(source.rglob("*")):
        if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png"}:
            # Keep the original extension in the stem to avoid x.png/x.jpg collisions.
            target = destination / path.relative_to(source).parent / (path.name + ".jpg")
            outputs.append(clean_image(path, target))
    return outputs


def placeholder(destination, label="FICTIONAL SAMPLE"):
    """A geometric placeholder, never a beneficiary portrait."""
    from PIL import ImageDraw
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (320, 240), "#e8eee7")
    draw = ImageDraw.Draw(img)
    draw.rectangle((24, 24, 296, 216), outline="#153f36", width=3)
    draw.text((65, 112), label, fill="#153f36")
    img.save(destination)
    return destination


def local_picture(image_root, filename):
    """Resolve a local picture inside the explicitly configured folder."""
    base = Path(image_root).resolve()
    path = (base / str(filename)).resolve()
    if base not in path.parents or not path.is_file():
        raise ValueError("Picture must be an existing file inside the configured image folder.")
    with Image.open(path) as opened:
        oriented = ImageOps.exif_transpose(opened).convert("RGB")
        oriented.thumbnail((400, 400))
        clean = Image.new("RGB", oriented.size)
        clean.paste(oriented)
        buffer = io.BytesIO()
        clean.save(buffer, format="JPEG", quality=82)
        buffer.seek(0)
    return buffer, clean.size


def register_pdf(rows, destination, title="Distribution register", variant="distribution", image_root=None):
    """Fit columns to landscape A4; repeat headings and paginate naturally."""
    columns = [("record_id", "Record ID"), ("name", "Name"), ("community", "Community"),
               ("household_size", "HH size"), ("card_number", "Card number")]
    if variant in {"distribution", "cash"}:
        columns.append(("amount", "Amount (demo units)"))
    elif variant != "validation":
        raise ValueError("Unknown register variant")
    columns.append(("", "Verification / signature"))
    if image_root is not None:
        columns.insert(-1, ("photo_filename", "Picture (sample)"))
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=11)
    width = landscape(A4)[0] - 64
    weights = [1, 1.7, 1.4, .65, 1.1] + ([1.1] if variant != "validation" else []) + [1.6]
    if image_root is not None:
        weights.insert(-1, 1)
    widths = [width * w / sum(weights) for w in weights]
    story = [Paragraph(escape(title), styles["Title"]),
             Paragraph("SYNTHETIC DEMONSTRATION — no real people or assistance allocations", styles["Normal"]), Spacer(1, 14)]
    for lga in sorted({r["lga"] for r in rows}):
        group = [r for r in rows if r["lga"] == lga]
        story += [Paragraph(escape(lga), styles["Heading2"])]
        matrix = [[Paragraph(escape(label), body) for _, label in columns]]
        for row in group:
            cells = []
            for key, _ in columns:
                if key == "photo_filename" and row.get(key):
                    picture, (w, h) = local_picture(image_root, row[key])
                    factor = min(46 / w, 46 / h)
                    cells.append(PDFImage(picture, width=w * factor, height=h * factor))
                else:
                    cells.append(Paragraph(escape(str(row.get(key, ""))), body))
            matrix.append(cells)
        table = Table(matrix, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eee7")),
                                  ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#ccd7ce")),
                                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("TOPPADDING", (0, 0), (-1, -1), 7),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
        story += [table, Spacer(1, 12)]
    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(32, 18, f"Seyi Abolarin | Synthetic workflow example | Page {doc.page}")
    SimpleDocTemplate(str(destination), pagesize=landscape(A4), leftMargin=32, rightMargin=32,
                      topMargin=28, bottomMargin=30, title=title, author="Seyi Abolarin").build(story, onFirstPage=footer, onLaterPages=footer)
    return destination


def register_docx(rows, destination):
    """Editable local register; no external office process or downloads."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.7), Inches(8.3)
    doc.add_heading("Synthetic distribution register", 0)
    doc.add_paragraph("Fictional records only. Review before any operational adaptation.")
    columns = ["record_id", "name", "community", "household_size", "card_number", "amount"]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light Shading Accent 1"
    for cell, key in zip(table.rows[0].cells, columns):
        cell.text = key.replace("_", " ").title()
    for row in rows:
        for cell, key in zip(table.add_row().cells, columns):
            cell.text = str(row.get(key, ""))
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)
    return destination


def prepare_local_attachments(source, destination):
    """Public replacement for authenticated picture downloads: local files only.

    Live URLs, tokens and remote attachment fetching are deliberately not included.
    Names are deterministic hashes and metadata is stripped from derived copies.
    """
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if source == destination or source in destination.parents or destination in source.parents:
        raise ValueError("Use separate, non-nested source and destination folders.")
    manifest = []
    for path in sorted(source.rglob("*")):
        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
            key = hashlib.sha256(path.relative_to(source).as_posix().encode()).hexdigest()[:20]
            result = clean_image(path, destination / f"sample-{key}.jpg")
            manifest.append({"attachment_id": key, "output": result.name})
    return manifest
